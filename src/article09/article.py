from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

from .config import ROOT, sha256_file, sha256_json
from .utils import atomic_write_json, task_complete, task_context


DOC_PYTHON = Path(
    r"C:\Users\Admin Riset\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe"
)
DOC_SKILL = Path(
    r"C:\Users\Admin Riset\.codex\plugins\cache\openai-primary-runtime\documents\26.1013.11914\skills\documents"
)


CITATION_PATTERNS = [
    (r"schul[d]? and killoran.*2019", "SCHULD-KILLORAN-2019"),
    (r"schuld et al.*2021", "SCHULD-SWEKE-MEYER-2021"),
    (r"havlicek.*2019", "HAVLICEK-2019"),
    (r"perez-salinas.*2020", "PEREZ-SALINAS-2020"),
    (r"mari et al.*2020", "MARI-2020"),
    (r"abbas et al.*2021", "ABBAS-2021"),
    (r"cerezo et al.*2021", "CEREZO-VQA-2021"),
    (r"cerezo et al.*2022", "CEREZO-CHALLENGES-2022"),
    (r"du et al.*2021", "DU-2021"),
    (r"caro et al.*2022", "CARO-2022"),
    (r"gil-fuster.*2024", "GIL-FUSTER-2024"),
    (r"oquab.*2024", "DINOV2-2024"),
    (r"huang.*2021", "HUANG-2021"),
    (r"egginger.*2024", "EGGINGER-2024"),
    (r"schnabel and roth.*2025", "SCHNABEL-ROTH-2025"),
    (r"mangini.*2022", "MANGINI-2022"),
    (r"wang.*2025", "WANG-2025"),
    (r"odagiu.*2025", "ODAGIU-2025"),
    (r"krizhevsky.*2009", "KRIZHEVSKY-2009"),
    (r"helber.*2019", "HELBER-2019"),
    (r"yang.*2023", "MEDMNIST-2023"),
    (r"goldberger.*2004", "GOLDBERGER-2004"),
    (r"demsar.*2006", "DEMSAR-2006"),
    (r"holm.*1979", "HOLM-1979"),
    (r"efron and tibshirani.*1993", "EFRON-TIBSHIRANI-1993"),
    (r"guo et al.*2017", "GUO-2017"),
    (r"gneiting and raftery.*2007", "GNEITING-RAFTERY-2007"),
    (r"thanasilp.*2024", "THANASILP-2024"),
]


def _set_font(run, name: str = "Times New Roman", size: float = 10.0, bold: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _set_cell_border(cell, *, top: tuple[str, str, str] | None = None, bottom: tuple[str, str, str] | None = None):
    """Apply restrained Springer-style horizontal rules and remove box borders."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        rule = top if edge == "top" else bottom if edge == "bottom" else None
        if rule is None:
            element.set(qn("w:val"), "nil")
        else:
            value, size, color = rule
            element.set(qn("w:val"), value)
            element.set(qn("w:sz"), size)
            element.set(qn("w:color"), color)


def _set_cell_margins(cell, top: int = 45, start: int = 55, bottom: int = 45, end: int = 55):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _format_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        row_properties.append(cant_split)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            if row_index == 0:
                _set_cell_border(cell, top=("single", "8", "000000"), bottom=("single", "6", "000000"))
            elif row_index == len(table.rows) - 1:
                _set_cell_border(cell, bottom=("single", "8", "000000"))
            else:
                _set_cell_border(cell, bottom=("single", "2", "D9D9D9"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.space_before = Pt(1)
                for run in paragraph.runs:
                    _set_font(run, size=8.5, bold=row_index == 0)
                    run.font.color.rgb = RGBColor(0, 0, 0)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]], caption: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(caption)
    _set_font(run, name="Arial Narrow", size=8.5, bold=True)
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            paragraph = cells[index].paragraphs[0]
            _append_text_with_citations(paragraph, str(value), size=8.5)
    _format_table(table)
    return table


def _add_omml_equation(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    node = OxmlElement("m:t")
    node.text = text
    run.append(node)
    math.append(run)
    math_para.append(math)
    paragraph._p.append(math_para)


def _append_reference_with_links(paragraph, reference: str):
    """Append a Springer-style reference with active DOI/URL hyperlinks."""
    cursor = 0
    for match in re.finditer(r"https?://\S+", reference):
        if match.start() > cursor:
            run = paragraph.add_run(reference[cursor : match.start()])
            _set_font(run, size=8.5)
        url = match.group(0).rstrip(".,;)")
        trailing = match.group(0)[len(url) :]
        _add_hyperlink(paragraph, url, url, size=8.5)
        if trailing:
            run = paragraph.add_run(trailing)
            _set_font(run, size=8.5)
        cursor = match.end()
    if cursor < len(reference):
        run = paragraph.add_run(reference[cursor:])
        _set_font(run, size=8.5)


def _format_reference_paragraph(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.widow_control = True


def _add_mendeley_bibliography(doc: Document, references: list[str]):
    """Create one active Mendeley bibliography field spanning separate entries."""
    if not references:
        return None
    paragraph = doc.add_paragraph()
    _format_reference_paragraph(paragraph)

    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run.append(begin)
    paragraph._p.append(begin_run)

    instruction_run = OxmlElement("w:r")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " ADDIN Mendeley Bibliography CSL_BIBLIOGRAPHY "
    instruction_run.append(instruction)
    paragraph._p.append(instruction_run)

    separate_run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run.append(separate)
    paragraph._p.append(separate_run)

    for index, reference in enumerate(references):
        if index:
            paragraph = doc.add_paragraph()
            _format_reference_paragraph(paragraph)
        _append_reference_with_links(paragraph, reference)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    paragraph._p.append(end_run)
    return paragraph


def _remove_paragraph_border(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    border = properties.find(qn("w:pBdr"))
    if border is not None:
        properties.remove(border)


def _add_page_number(section):
    """Add the automatic Word PAGE field required by the target journal."""
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    instruction_run = OxmlElement("w:r")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    instruction_run.append(instruction)
    separate_run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run.append(separate)
    display_run = OxmlElement("w:r")
    display_text = OxmlElement("w:t")
    display_text.text = "1"
    display_run.append(display_text)
    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    paragraph._p.extend([begin_run, instruction_run, separate_run, display_run, end_run])


def _set_columns(section, count: int):
    sect_pr = section._sectPr
    columns = sect_pr.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        sect_pr.append(columns)
    columns.set(qn("w:num"), str(count))
    columns.set(qn("w:space"), str(300 if count == 2 else 0))


def _configure_section(section, columns: int = 1):
    # Measured from the supplied published QMI reference: 595.276 x 790.866 pt.
    section.page_width = Inches(595.276 / 72)
    section.page_height = Inches(790.866 / 72)
    section.top_margin = Inches(0.47)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(51.024 / 72)
    section.right_margin = Inches(51.024 / 72)
    _set_columns(section, columns)


def _switch_columns(doc: Document, columns: int):
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    _configure_section(section, columns)
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True
    return section


def _add_full_width_table(doc: Document, headers: list[str], rows: list[list[str]], caption: str):
    _switch_columns(doc, 1)
    table = _add_table(doc, headers, rows, caption)
    _switch_columns(doc, 2)
    return table


def _add_full_width_figure(doc: Document, stem: str, caption: str, width: float = 6.65):
    _switch_columns(doc, 1)
    _add_figure(doc, stem, caption, width=width)
    _switch_columns(doc, 2)


def _add_hyperlink(paragraph, text: str, url: str, size: float = 9.0):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(int(size * 2)))
    font_size_cs = OxmlElement("w:szCs")
    font_size_cs.set(qn("w:val"), str(int(size * 2)))
    properties.extend([color, underline, font_size, font_size_cs])
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _citation_ids(shown: str) -> list[str]:
    normalized = shown.lower().replace("–", "-")
    found = []
    for pattern, citation_id in CITATION_PATTERNS:
        if re.search(pattern, normalized) and citation_id not in found:
            found.append(citation_id)
    return found


def _add_mendeley_citation(paragraph, shown: str, ids: list[str], name: str, size: float):
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run.append(begin)
    paragraph._p.append(begin_run)

    instruction_run = OxmlElement("w:r")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    payload = {
        "citationItems": [{"id": item} for item in ids],
        "properties": {"formattedCitation": shown, "plainCitation": shown, "noteIndex": 0},
        "schema": "https://github.com/citation-style-language/schema/raw/master/csl-citation.json",
    }
    instruction.text = " ADDIN CSL_CITATION " + json.dumps(payload, separators=(",", ":")) + " "
    instruction_run.append(instruction)
    paragraph._p.append(instruction_run)

    separate_run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run.append(separate)
    paragraph._p.append(separate_run)

    visible = paragraph.add_run(shown)
    _set_font(visible, name=name, size=size)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    paragraph._p.append(end_run)


def _append_text_with_citations(paragraph, text: str, name: str = "Times New Roman", size: float = 10.0):
    cursor = 0
    for match in re.finditer(r"\([^()]*?(?:19|20)\d{2}[^()]*?\)", text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            _set_font(run, name=name, size=size)
        shown = match.group(0)
        ids = _citation_ids(shown)
        if ids:
            _add_mendeley_citation(paragraph, shown, ids, name, size)
        else:
            run = paragraph.add_run(shown)
            _set_font(run, name=name, size=size)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        _set_font(run, name=name, size=size)


def _add_figure(doc: Document, stem: str, caption: str, width: float = 6.5):
    path = ROOT / "results" / "figures" / f"{stem}.png"
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = paragraph.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("title", caption)
    picture._inline.docPr.set("descr", caption)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.keep_together = True
    caption_paragraph.paragraph_format.space_after = Pt(4)
    match = re.match(r"(Figure\s+\d+)\s+(.*)", caption)
    if match:
        label = caption_paragraph.add_run(match.group(1) + " ")
        _set_font(label, name="Arial Narrow", size=8.5, bold=True)
        body = caption_paragraph.add_run(match.group(2))
        _set_font(body, size=8.5)
    else:
        run = caption_paragraph.add_run(caption)
        _set_font(run, size=8.5)


def _paragraph(
    doc: Document,
    text: str,
    bold_lead: str | None = None,
    *,
    first_line_indent: bool = True,
):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.25 if first_line_indent else 0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(3)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        _set_font(lead, bold=True)
        _append_text_with_citations(paragraph, text[len(bold_lead) :])
    else:
        _append_text_with_citations(paragraph, text)
    return paragraph


def _result_narrative() -> dict:
    test = pd.read_csv(ROOT / "results" / "metrics" / "all_test_metrics.csv")
    validation = pd.read_csv(ROOT / "results" / "metrics" / "all_validation_metrics.csv")
    compression = pd.read_csv(ROOT / "results" / "statistics" / "compression_penalty.csv")
    gaps = pd.read_csv(ROOT / "results" / "statistics" / "head_gaps.csv")
    friedman = pd.read_csv(ROOT / "results" / "statistics" / "friedman.csv")
    main = test[test.reducer != "full384"]
    grouped = main.groupby(["dataset", "reducer", "head"], as_index=False).agg(
        macro_f1_mean=("macro_f1", "mean"), macro_f1_sd=("macro_f1", "std"), accuracy_mean=("accuracy", "mean")
    )
    validation_grouped = (
        validation[validation["reducer"] != "full384"]
        .groupby(["dataset", "reducer", "head"], as_index=False)
        .agg(validation_macro_f1=("macro_f1", "mean"))
    )
    selected = (
        validation_grouped.sort_values(
            ["dataset", "validation_macro_f1"], ascending=[True, False]
        )
        .groupby("dataset")
        .head(1)
    )
    winners = selected.merge(
        grouped, on=["dataset", "reducer", "head"], how="left", validate="one_to_one"
    )
    selected_vqc = (
        validation_grouped[validation_grouped["head"] == "vqc"]
        .sort_values(["dataset", "validation_macro_f1"], ascending=[True, False])
        .groupby("dataset")
        .head(1)
    )
    vqc = selected_vqc.merge(
        grouped, on=["dataset", "reducer", "head"], how="left", validate="one_to_one"
    )
    compression_abs = float(compression.compression_penalty_mean.abs().mean())
    matched = gaps[gaps.comparison == "vqc_minus_mlp"]
    head_abs = float(matched.mean_difference.abs().mean())
    return {
        "grouped": grouped,
        "winners": winners,
        "vqc": vqc,
        "compression_abs": compression_abs,
        "head_abs": head_abs,
        "compression": compression,
        "gaps": gaps,
        "friedman": friedman,
    }


REFERENCES = [
    "Abbas A, Sutter D, Zoufal C, Lucchi A, Figalli A, Woerner S (2021) The power of quantum neural networks. Nat Comput Sci 1:403-409. https://doi.org/10.1038/s43588-021-00084-1",
    "Baldi P, Hornik K (1989) Neural networks and principal component analysis. Neural Netw 2(1):53-58. https://doi.org/10.1016/0893-6080(89)90014-2",
    "Caro MC, Huang HY, Cerezo M, Sharma K, Sornborger A, Cincio L, Coles PJ (2022) Generalization in quantum machine learning from few training data. Nat Commun 13:4919. https://doi.org/10.1038/s41467-022-32550-3",
    "Cerezo M et al (2021) Variational quantum algorithms. Nat Rev Phys 3:625-644. https://doi.org/10.1038/s42254-021-00348-9",
    "Cerezo M, Verdon G, Huang HY, Cincio L, Coles PJ (2022) Challenges and opportunities in quantum machine learning. Nat Comput Sci 2:567-576. https://doi.org/10.1038/s43588-022-00311-3",
    "Demsar J (2006) Statistical comparisons of classifiers over multiple data sets. J Mach Learn Res 7:1-30",
    "Du Y, Hsieh MH, Liu T, You S, Tao D (2021) Learnability of quantum neural networks. PRX Quantum 2:040337. https://doi.org/10.1103/PRXQuantum.2.040337",
    "Efron B, Tibshirani RJ (1993) An Introduction to the Bootstrap. Chapman and Hall, New York",
    "Egginger S, Sakhnenko A, Lorenz JM (2024) A hyperparameter study for quantum kernel methods. Quantum Mach Intell 6:44. https://doi.org/10.1007/s42484-024-00172-1",
    "Gil-Fuster E, Eisert J, Bravo-Prieto C (2024) Understanding quantum machine learning also requires rethinking generalization. Nat Commun 15:2277. https://doi.org/10.1038/s41467-024-45882-z",
    "Gneiting T, Raftery AE (2007) Strictly proper scoring rules, prediction, and estimation. J Am Stat Assoc 102(477):359-378. https://doi.org/10.1198/016214506000001437",
    "Goldberger J, Hinton GE, Roweis S, Salakhutdinov RR (2004) Neighbourhood components analysis. Adv Neural Inf Process Syst 17",
    "Guo C, Pleiss G, Sun Y, Weinberger KQ (2017) On calibration of modern neural networks. Proceedings of the 34th International Conference on Machine Learning, pp 1321-1330",
    "Havlicek V et al (2019) Supervised learning with quantum-enhanced feature spaces. Nature 567:209-212. https://doi.org/10.1038/s41586-019-0980-2",
    "Helber P, Bischke B, Dengel A, Borth D (2019) EuroSAT: A novel dataset and deep learning benchmark for land use and land cover classification. IEEE J Sel Top Appl Earth Obs Remote Sens 12(7):2217-2226. https://doi.org/10.1109/JSTARS.2019.2918242",
    "Holm S (1979) A simple sequentially rejective multiple test procedure. Scand J Stat 6(2):65-70",
    "Huang HY, Broughton M, Mohseni M, Babbush R, Boixo S, Neven H, McClean JR (2021) Power of data in quantum machine learning. Nat Commun 12:2631. https://doi.org/10.1038/s41467-021-22539-9",
    "Krizhevsky A (2009) Learning multiple layers of features from tiny images. Technical report, University of Toronto",
    "Mangini S, Marruzzo A, Piantanida M, Gerace D, Bajoni D, Macchiavello C (2022) Quantum neural network autoencoder and classifier applied to an industrial case study. Quantum Mach Intell 4:13. https://doi.org/10.1007/s42484-022-00070-4",
    "Mari A, Bromley TR, Izaac J, Schuld M, Killoran N (2020) Transfer learning in hybrid classical-quantum neural networks. Quantum 4:340. https://doi.org/10.22331/q-2020-10-09-340",
    "Odagiu P et al (2025) Learning reduced representations for quantum classifiers. Quantum Mach Intell 7:113. https://doi.org/10.1007/s42484-025-00331-y",
    "Oquab M et al (2024) DINOv2: Learning robust visual features without supervision. Trans Mach Learn Res",
    "Perez-Salinas A, Cervera-Lierta A, Gil-Fuster E, Latorre JI (2020) Data re-uploading for a universal quantum classifier. Quantum 4:226. https://doi.org/10.22331/q-2020-02-06-226",
    "Schnabel J, Roth M (2025) Quantum kernel methods under scrutiny: a benchmarking study. Quantum Mach Intell 7:58. https://doi.org/10.1007/s42484-025-00273-5",
    "Schuld M, Killoran N (2019) Quantum machine learning in feature Hilbert spaces. Phys Rev Lett 122:040504. https://doi.org/10.1103/PhysRevLett.122.040504",
    "Schuld M, Sweke R, Meyer JJ (2021) Effect of data encoding on the expressive power of variational quantum-machine-learning models. Phys Rev A 103:032430. https://doi.org/10.1103/PhysRevA.103.032430",
    "Thanasilp S, Wang S, Cerezo M, Holmes Z (2024) Exponential concentration in quantum kernel methods. Nat Commun 15:5200. https://doi.org/10.1038/s41467-024-49287-w",
    "Wang P, Myers CR, Hollenberg LCL, Parampalli U (2025) Quantum Hamiltonian embedding of images for data reuploading classifiers. Quantum Mach Intell 7:35. https://doi.org/10.1007/s42484-025-00247-7",
    "Yang J et al (2023) MedMNIST v2: A large-scale lightweight benchmark for 2D and 3D biomedical image classification. Sci Data 10:41. https://doi.org/10.1038/s41597-022-01721-8",
]


def build_manuscript(output: Path) -> None:
    result = _result_narrative()
    doc = Document()
    update_fields = doc.settings.element.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        doc.settings.element.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    # The published QMI layout relies on automatic hyphenation to keep fully
    # justified two-column text from developing conspicuous word gaps.
    settings = doc.settings.element
    auto_hyphenation = settings.find(qn("w:autoHyphenation"))
    if auto_hyphenation is None:
        auto_hyphenation = OxmlElement("w:autoHyphenation")
        settings.append(auto_hyphenation)
    auto_hyphenation.set(qn("w:val"), "true")
    hyphenation_zone = settings.find(qn("w:hyphenationZone"))
    if hyphenation_zone is None:
        hyphenation_zone = OxmlElement("w:hyphenationZone")
        settings.append(hyphenation_zone)
    hyphenation_zone.set(qn("w:val"), "180")
    consecutive_hyphens = settings.find(qn("w:consecutiveHyphenLimit"))
    if consecutive_hyphens is None:
        consecutive_hyphens = OxmlElement("w:consecutiveHyphenLimit")
        settings.append(consecutive_hyphens)
    consecutive_hyphens.set(qn("w:val"), "2")
    doc.core_properties.title = (
        "Disentangling Representation Compression and Quantum-Head Effects in Multi-Domain Quantum Transfer Learning"
    )
    doc.core_properties.author = "Fajri Yanuar Shiddiq Juanda; Kiki Ahmad Baihaqi"
    doc.core_properties.subject = "Quantum machine learning; representation compression; transfer learning"
    section = doc.sections[0]
    _configure_section(section, columns=1)
    _add_page_number(section)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    styles["Normal"].font.size = Pt(10)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Arial Narrow"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial Narrow")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Narrow")
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Title"].font.size = Pt(16)
    styles["Heading 1"].font.size = Pt(11)
    styles["Heading 2"].font.size = Pt(10)
    styles["Heading 3"].font.size = Pt(10)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[style_name].paragraph_format.keep_with_next = True
        styles[style_name].paragraph_format.space_before = Pt(7)
        styles[style_name].paragraph_format.space_after = Pt(3)

    title = doc.add_paragraph(style="Title")
    _remove_paragraph_border(title)
    title_style_properties = styles["Title"]._element.get_or_add_pPr()
    title_style_border = title_style_properties.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_properties.remove(title_style_border)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run(
        "Disentangling Representation Compression and Quantum-Head Effects in Multi-Domain Quantum Transfer Learning"
    )
    _set_font(run, name="Arial Narrow", size=16, bold=True)
    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = authors.add_run("Fajri Yanuar Shiddiq Juanda")
    _set_font(run, name="Arial Narrow", size=10, bold=True)
    run = authors.add_run("1")
    _set_font(run, name="Arial Narrow", size=7, bold=True)
    run.font.superscript = True
    run = authors.add_run("*  ·  Kiki Ahmad Baihaqi")
    _set_font(run, name="Arial Narrow", size=10, bold=True)
    run = authors.add_run("1")
    _set_font(run, name="Arial Narrow", size=7, bold=True)
    run.font.superscript = True
    author_orcid = doc.add_paragraph()
    author_orcid.paragraph_format.space_after = Pt(1)
    lead = author_orcid.add_run("Fajri Yanuar Shiddiq Juanda — ORCID: ")
    _set_font(lead, size=9)
    _add_hyperlink(author_orcid, "0009-0007-8894-0265", "https://orcid.org/0009-0007-8894-0265")
    author_orcid.add_run("\n")
    lead = author_orcid.add_run("Kiki Ahmad Baihaqi — ORCID: ")
    _set_font(lead, size=9)
    _add_hyperlink(author_orcid, "0000-0002-2430-9235", "https://orcid.org/0000-0002-2430-9235")
    affiliation = doc.add_paragraph()
    affiliation.paragraph_format.space_after = Pt(1)
    run = affiliation.add_run(
        "1 Department of Informatics Engineering, Faculty of Computer Science, Universitas Buana Perjuangan Karawang, Karawang, Indonesia"
    )
    _set_font(run, size=9)
    contacts = doc.add_paragraph()
    contacts.paragraph_format.space_after = Pt(6)
    run = contacts.add_run(
        "Email: fajriyanuar1@gmail.com; kikiahmad@ubpkarawang.ac.id\n"
        "*Corresponding author: Fajri Yanuar Shiddiq Juanda, fajriyanuar1@gmail.com"
    )
    _set_font(run, size=9)

    doc.add_heading("Abstract", level=1)
    winners = result["winners"]
    best_text = "; ".join(
        f"{row.dataset}: {row.reducer.upper()} + {row.head.upper()} macro-F1 {row.macro_f1_mean:.3f} ± {row.macro_f1_sd:.3f}"
        for row in winners.itertuples()
    )
    relation = (
        "the average absolute compression penalty exceeded the matched VQC-MLP head gap"
        if result["compression_abs"] > result["head_abs"]
        else "the matched head gap was at least as large as the average absolute compression penalty"
    )
    abstract = (
        "Hybrid quantum visual classifiers compress high-dimensional classical representations to a small number of circuit inputs, "
        "but conventional comparisons rarely separate information lost before the circuit from differences introduced by the decision head. "
        "This study used frozen DINOv2 ViT-S/14 embeddings from CIFAR-10, EuroSAT RGB, and PathMNIST and crossed four six-dimensional reducers "
        "with logistic regression, RBF-SVM, an exactly parameter-matched multilayer perceptron, and a six-qubit variational quantum classifier. "
        "The leakage-aware protocol used duplicate-group isolation, a fixed 150 examples per class, ten matched seeds, validation-only selection, "
        "and a locked test set. The experiment completed 480 reducer-head runs plus 60 full-space references. "
        f"The validation-selected best configurations were {best_text}. Across evaluated conditions, {relation}. "
        "These results locate the dominant bottleneck empirically rather than attributing every downstream difference to the quantum circuit. "
        "The findings are limited to analytic statevector simulation and do not constitute evidence of quantum advantage or speedup."
    )
    _paragraph(doc, abstract, first_line_indent=False)
    keywords = doc.add_paragraph()
    lead = keywords.add_run("Keywords: ")
    _set_font(lead, bold=True)
    tail = keywords.add_run(
        "quantum machine learning; DINOv2; representation compression; variational quantum classifier; transfer learning; calibration"
    )
    _set_font(tail)

    _switch_columns(doc, 2)

    doc.add_heading("1 Introduction", level=1)
    _paragraph(
        doc,
        "Quantum transfer learning typically places a low-qubit variational circuit after a classical visual encoder. The apparent simplicity hides two distinct sources of error: the classical-to-quantum bottleneck can discard discriminative structure, while the decision head can fail to exploit what remains. Comparing a high-dimensional classical baseline with a compressed quantum model therefore confounds representation quality and classifier behavior.",
    )
    _paragraph(
        doc,
        "Recent work has established quantum feature spaces, variational learning, data re-uploading, and hybrid transfer-learning architectures (Schuld and Killoran 2019; Havlicek et al 2019; Perez-Salinas et al 2020; Mari et al 2020; Abbas et al 2021; Cerezo et al 2021, 2022). Encoding design can change the function family available to a variational model (Schuld et al 2021), while trainability and finite-data generalization remain separate concerns (Du et al 2021; Caro et al 2022; Gil-Fuster et al 2024). In parallel, self-supervised encoders have made frozen representations increasingly transferable (Oquab et al 2024), making it possible to control the encoder while studying the downstream bottleneck. However, a fair comparison must give classical and quantum heads the same samples, splits, six-dimensional vectors, selection criterion, and seed schedule.",
    )
    _paragraph(
        doc,
        "Empirical QML studies increasingly emphasize strong classical baselines and explicit accounting for preprocessing. Classical access to data can make conventional learners highly competitive even when quantum representations appear difficult to simulate (Huang et al 2021). QMI benchmarks likewise show that hyperparameters and classical components can dominate observed behavior (Egginger et al 2024; Schnabel and Roth 2025), while application studies illustrate both the promise and the limits of compressed quantum classifiers (Mangini et al 2022; Wang et al 2025). Most directly, recent work on learned reduced representations confirms that dimensionality reduction deserves treatment as an experimental factor rather than an incidental preprocessing step (Odagiu et al 2025).",
    )
    _paragraph(
        doc,
        "The present study differs from prior reduction-focused QML work in three respects. First, every reducer is crossed with four predictive heads, so the representation and classifier are estimated as separate experimental factors. Second, the VQC is compared with a parameter-matched MLP on identical vectors rather than only with an unconstrained high-dimensional baseline. Third, the analysis is repeated across three visual domains with paired ten-seed inference, full-space references, and calibration diagnostics. This design turns the usual end-to-end leaderboard into an attribution study of where performance is lost.",
    )
    _paragraph(
        doc,
        "This study asks whether compression or the quantum head is the larger bottleneck. Its contribution is a controlled reducer-by-head factorial benchmark across general-object, remote-sensing, and histopathology domains. The design combines a frozen 384D DINOv2 representation, four reducers at an identical 6D budget, exact parameter parity between the MLP and VQC, full-space references, representation-geometry diagnostics, calibration metrics, and paired multi-seed statistics.",
    )
    _add_full_width_table(
        doc,
        ["Research stream", "Typical design", "Unresolved issue", "Control in this study"],
        [
            ["Quantum feature spaces and encodings (Schuld and Killoran 2019; Schuld et al 2021)", "Quantum kernel or circuit", "Representation quality may vary", "Identical six-dimensional input"],
            ["Hybrid transfer learning (Mari et al 2020; Mangini et al 2022)", "Classical encoder plus quantum head", "Encoder and head effects can be confounded", "Frozen DINOv2 for every condition"],
            ["Reduced quantum representations (Odagiu et al 2025)", "Reducer followed by a quantum classifier", "Reducer and classifier are rarely crossed factorially", "Four reducers crossed with four heads"],
            ["QML benchmarking (Egginger et al 2024; Schnabel and Roth 2025)", "Hyperparameter-sensitive comparison", "Classical components may dominate", "Matched data, selection, seeds, and capacity"],
            ["Model calibration (Guo et al 2017; Gneiting and Raftery 2007)", "Confidence quality after training", "Often omitted in QML comparisons", "ECE, NLL, and Brier score"],
        ],
        "Table 1 Related methodological streams and the controls used in this study",
    )

    doc.add_heading("2 Materials and Methods", level=1)
    doc.add_heading("2.1 Datasets and leakage audit", level=2)
    _paragraph(
        doc,
        "CIFAR-10 (Krizhevsky 2009), EuroSAT RGB (Helber et al 2019), and PathMNIST (Yang et al 2023) provide ten, ten, and nine classes. Official test data were retained for CIFAR-10; PathMNIST used its official train-validation-test split; EuroSAT used a deterministic group-aware 70:15:15 split. Before freezing the partitions, every image was decoded, hashed at pixel level, screened by perceptual hash, and verified using structural similarity. Exact or confirmed near-duplicate groups were indivisible. Any group crossing an official split or carrying inconsistent labels was quarantined.",
    )
    audit = pd.read_csv(ROOT / "results" / "tables" / "dataset_audit_summary.csv")
    audit_rows = []
    for _, row in audit.iterrows():
        audit_rows.append(
            [
                str(row.get("Dataset", "")),
                str(row.get("Total decoded", "")),
                str(row.get("Duplicate groups", "")),
                str(row.get("Quarantined", "")),
                str(row.get("Train split", "")),
                str(row.get("Val split", "")),
                str(row.get("Test split", "")),
            ]
        )
    _add_full_width_table(
        doc,
        ["Dataset", "Decoded", "Duplicate groups", "Quarantined", "Train", "Validation", "Test"],
        audit_rows,
        "Table 2 Dataset audit and frozen split counts",
    )
    _add_full_width_figure(doc, "figure_02_dataset_samples", "Figure 1 Representative samples generated by the reproducible analysis notebook")

    doc.add_heading("2.2 Frozen representation and compression", level=2)
    _paragraph(
        doc,
        "Every image was converted to RGB, resized deterministically to 224 by 224 pixels using bicubic interpolation, normalized with ImageNet statistics, and processed once by frozen DINOv2 ViT-S/14 (Oquab et al 2024). The 384D normalized CLS token was cached as float32. No image augmentation or encoder fine-tuning was used, ensuring that the only experimental factor before classification was the reducer.",
    )
    _paragraph(
        doc,
        "PCA-6 represented unsupervised variance preservation; LDA-6 represented supervised linear discrimination; NCA-6 represented supervised metric learning (Goldberger et al 2004); and SLP-6 used a learned 384-to-6 tanh bottleneck whose temporary classification layer was discarded. Each reducer was fitted on training data only. A second training-fitted standardizer mapped every reduced coordinate to an angle-compatible range.",
    )
    _add_omml_equation(doc, "x̃ⱼ = π tanh(zⱼ)")
    _add_full_width_figure(doc, "figure_01_pipeline", "Figure 2 Controlled reducer-by-head factorial design; editable source prepared in draw.io")

    doc.add_heading("2.3 Classical and quantum heads", level=2)
    _paragraph(
        doc,
        "All four heads received exactly the same scaled six-dimensional vectors. Logistic regression and RBF-SVM provided linear and nonlinear classical references. The MLP comprised a bias-free 6-to-6 hidden layer, tanh activation, and a 6-to-C output layer. The VQC used six qubits and two data-reuploading layers; each layer applied RY input encoding, one trainable Rot gate per qubit, and ring CNOT entanglement before six Pauli-Z expectation values were passed to a 6-to-C linear output.",
    )
    _add_omml_equation(doc, "Pᵥᵩ꜀ = 6 × 3 × 2 + 6C + C = 36 + 7C")
    _paragraph(
        doc,
        "The MLP parameter count is 6 x 6 + 6C + C, exactly matching the VQC at 106 trainable parameters for ten-class datasets and 99 for PathMNIST. This parity prevents model capacity from becoming an uncontrolled explanation for a classical-quantum difference.",
    )

    doc.add_heading("2.4 Selection, test isolation, and statistics", level=2)
    _paragraph(
        doc,
        "Each final seed selected a stratified 150 examples per class from training data, while a single balanced validation set contained 100 examples per class. Hyperparameters were selected with pilot seed 2026 by validation macro-F1, with validation negative log-likelihood as tie-breaker. The selected settings and data hashes were frozen before any test-label access. Final seeds were 11, 22, 33, 44, 55, 66, 77, 88, 99, and 111.",
    )
    _paragraph(
        doc,
        "Macro-F1 was the primary endpoint. Secondary endpoints were balanced accuracy, accuracy, MCC, macro precision and recall, one-vs-rest macro ROC-AUC, NLL, expected calibration error, multiclass Brier score, fit time, and inference time. Reducer effects were tested with Friedman tests followed by planned Wilcoxon signed-rank comparisons and Holm correction (Demsar 2006; Holm 1979). Paired model gaps and compression penalties used 10,000 bootstrap resamples over seed differences (Efron and Tibshirani 1993). Datasets were summarized separately rather than pooled as pseudo-independent replicates.",
    )
    _add_omml_equation(doc, "Δcompression(r) = F1RBF-384 - F1RBF-6,r")
    _add_omml_equation(doc, "ΔVQC-MLP(r) = F1VQC,r - F1MLP,r")

    doc.add_heading("3 Results", level=1)
    _paragraph(
        doc,
        "All reported test values were produced after validation-only selection and protocol freezing. Table 3 summarizes the mean and standard deviation across the ten matched seeds. Figure 3 exposes interactions between reducer and head that would be hidden by a single best-model comparison.",
    )
    table_rows = []
    for row in result["grouped"].itertuples():
        table_rows.append(
            [row.dataset, row.reducer, row.head, f"{row.macro_f1_mean:.4f} +/- {row.macro_f1_sd:.4f}", f"{row.accuracy_mean:.4f}"]
        )
    _add_full_width_table(
        doc,
        ["Dataset", "Reducer", "Head", "Macro-F1 mean +/- SD", "Accuracy mean"],
        table_rows,
        "Table 3 Test performance across ten matched seeds",
    )
    _add_full_width_figure(doc, "figure_03_macro_f1_heatmap", "Figure 3 Mean test macro-F1 for each controlled reducer-head condition")

    doc.add_heading("3.1 Compression and head effects", level=2)
    relation_sentence = (
        f"The mean absolute compression penalty was {result['compression_abs']:.4f}, exceeding the mean absolute matched VQC-MLP gap of {result['head_abs']:.4f}."
        if result["compression_abs"] > result["head_abs"]
        else f"The mean absolute matched VQC-MLP gap was {result['head_abs']:.4f}, at least as large as the mean absolute compression penalty of {result['compression_abs']:.4f}."
    )
    _paragraph(doc, relation_sentence)
    compression_rows = [
        [row.dataset, row.reducer, f"{row.rbf384_mean:.4f}", f"{row.rbf6_mean:.4f}", f"{row.compression_penalty_mean:.4f}", f"[{row.ci_low:.4f}, {row.ci_high:.4f}]"]
        for row in result["compression"].itertuples()
    ]
    _add_full_width_table(
        doc,
        ["Dataset", "Reducer", "RBF 384D", "RBF 6D", "Penalty", "95% CI"],
        compression_rows,
        "Table 4 Compression penalty relative to the full-space RBF reference",
    )
    matched_gaps = result["gaps"][result["gaps"]["comparison"] == "vqc_minus_mlp"]
    gap_rows = [
        [
            row.dataset,
            row.reducer,
            f"{row.mean_difference:.4f}",
            f"[{row.ci_low:.4f}, {row.ci_high:.4f}]",
            f"{row.holm_p:.4g}",
        ]
        for row in matched_gaps.itertuples()
    ]
    _add_full_width_table(
        doc,
        ["Dataset", "Reducer", "VQC minus MLP", "95% paired bootstrap CI", "Holm adjusted p"],
        gap_rows,
        "Table 5 Matched quantum-classical head differences across ten paired seeds",
    )
    friedman_rows = [
        [row.dataset, f"{row.statistic:.3f}", f"{row.p_value:.3g}", str(int(row.n_seeds))]
        for row in result["friedman"].itertuples()
    ]
    _add_full_width_table(
        doc,
        ["Dataset", "Friedman statistic", "p value", "Paired seed blocks"],
        friedman_rows,
        "Table 6 Friedman tests of VQC performance across the four reducers",
    )
    _paragraph(
        doc,
        "Reducer choice affected VQC macro-F1 in all three domains (Friedman statistic 27.12, p = 5.56e-6 for each dataset). "
        "Every VQC-minus-MLP contrast was negative, its paired bootstrap confidence interval excluded zero, and its Holm-adjusted Wilcoxon p value was 0.0156. "
        "The smallest disadvantage occurred with SLP-6, indicating that the learned supervised bottleneck mitigated, but did not remove, the quantum-head deficit.",
    )
    _add_full_width_figure(doc, "figure_04_compression_and_head_gaps", "Figure 4 Compression penalties and matched quantum-head gaps")

    doc.add_heading("3.2 Representation geometry and error structure", level=2)
    _paragraph(
        doc,
        "Distance preservation, neighborhood trustworthiness, Fisher separation, and silhouette score were computed on the fixed validation subset. Their association with VQC macro-F1 indicates whether circuit behavior tracks the geometry delivered by the reducer rather than the reducer label alone.",
    )
    _add_full_width_figure(doc, "figure_05_representation_geometry", "Figure 5 Representation geometry in relation to VQC performance")
    _add_full_width_figure(
        doc,
        "figure_06a_confusion_cifar10",
        "Figure 6a Row-normalized CIFAR-10 test confusion matrix for validation-selected SLP-6 with logistic regression",
        width=5.45,
    )
    _add_full_width_figure(
        doc,
        "figure_06b_confusion_eurosat",
        "Figure 6b Row-normalized EuroSAT test confusion matrix for validation-selected SLP-6 with logistic regression",
        width=5.45,
    )
    _add_full_width_figure(
        doc,
        "figure_06c_confusion_pathmnist",
        "Figure 6c Row-normalized PathMNIST test confusion matrix for validation-selected SLP-6 with RBF-SVM",
        width=5.45,
    )
    _add_full_width_figure(doc, "figure_07_calibration", "Figure 7 Reliability diagrams for validation-selected configurations")

    doc.add_heading("3.3 Runtime", level=2)
    _paragraph(
        doc,
        "Runtime was recorded for the observed software implementation. Because the VQC used analytic statevector simulation on classical hardware, these measurements describe implementation cost only and cannot support a quantum-speedup claim.",
    )
    _add_full_width_figure(doc, "figure_08_runtime", "Figure 8 Mean observed fit time by head and dataset")

    doc.add_heading("4 Discussion and conclusion", level=1)
    _paragraph(
        doc,
        "The controlled design changes the interpretation of a classical-quantum performance gap. When the full-space reference declines substantially after six-dimensional projection, the downstream model is operating after measurable information loss. Conversely, a persistent VQC-MLP gap on identical vectors isolates the decision head more directly than conventional transfer-learning comparisons.",
    )
    _paragraph(
        doc,
        "Supervised reducers can favor class separation under a severe information budget, but their validation selection must remain separate from test evaluation. The ten-seed paired design also shows whether a favorable mean is stable or depends on a small number of training subsets. Calibration provides a complementary view because similar macro-F1 values can coexist with materially different confidence quality.",
    )
    _paragraph(
        doc,
        "The results should be read as a benchmark of low-dimensional predictive heads, not as evidence that a quantum circuit offers computational advantage. This interpretation is consistent with evidence that optimized classical learners can remain competitive when data are available (Huang et al 2021) and with recent QMI benchmarks that find no broad quantum-model advantage across conventional tasks (Schnabel and Roth 2025). A competitive VQC result would show that the chosen circuit can exploit a representation already shaped by classical learning; it would not show speedup, hardware advantage, or universal superiority. The circuit uses a shallow data-reuploading design to avoid unnecessary depth, while the analytic statevector setting deliberately separates representation and head effects from shot noise or device noise (Perez-Salinas et al 2020; Thanasilp et al 2024).",
    )

    _paragraph(
        doc,
        "Several boundaries constrain the interpretation. The VQC was evaluated with analytic statevector simulation; finite-shot sampling, device noise, connectivity constraints, transpilation overhead, and physical-hardware drift were outside the scope. DINOv2 was frozen, so the conclusions concern downstream bottlenecks rather than end-to-end optimization. The benchmark used three image domains and a fixed six-dimensional budget, and other encoders, dimensions, and circuit families may produce different interactions.",
    )
    _paragraph(
        doc,
        "In conclusion, this study separates pre-quantum representation loss from quantum-head behavior by crossing four reducers and four matched heads on identical six-dimensional vectors from a frozen self-supervised encoder. Full-space references quantify information lost before the circuit, while parameter parity and paired seeds isolate the downstream comparison. The evidence identifies where the predictive bottleneck occurs under the evaluated protocol without supporting a claim of quantum advantage.",
    )

    _paragraph(
        doc,
        "Supplementary Information. The supplementary material accompanying this submission contains the executable notebook with stored outputs, configuration files, split and artifact hashes, prediction-level outputs, model checkpoints, environment records, and editable figure sources.",
        bold_lead="Supplementary Information.",
        first_line_indent=False,
    )
    _paragraph(
        doc,
        "Acknowledgements. The authors have no acknowledgements to declare.",
        bold_lead="Acknowledgements.",
        first_line_indent=False,
    )
    _paragraph(
        doc,
        "Author Contributions. F.Y.S.J. contributed conceptualization, methodology, software, formal analysis, investigation, data curation, visualization, writing—original draft, and project administration. K.A.B. contributed supervision, validation, resources, and writing—review and editing. Both authors reviewed and approved the final manuscript.",
        bold_lead="Author Contributions.",
        first_line_indent=False,
    )
    _paragraph(
        doc,
        "Funding. The authors received no specific funding for this work.",
        bold_lead="Funding.",
        first_line_indent=False,
    )
    _paragraph(
        doc,
        "Data Availability. CIFAR-10, EuroSAT RGB, and PathMNIST are publicly available from the original repositories cited in the Methods. The exact split manifests, sample identifiers, audit reports, hashes, derived metrics, and prediction-level outputs used in this study are included in the supplementary reproducibility package and are available from the corresponding author.",
        bold_lead="Data Availability.",
        first_line_indent=False,
    )
    _paragraph(
        doc,
        "Code Availability. The executable notebook, source code, configuration files, environment records, and trained artifacts are included in the supplementary reproducibility package. These materials will also be archived in a public repository upon acceptance.",
        bold_lead="Code Availability.",
        first_line_indent=False,
    )

    doc.add_heading("Declarations", level=1)
    _paragraph(
        doc,
        "Competing Interests. The authors declare that they have no competing interests.",
        bold_lead="Competing Interests.",
        first_line_indent=False,
    )
    _paragraph(
        doc,
        "Open Access. The authors request open-access publication under the Creative Commons Attribution 4.0 International License, subject to the journal's publication agreement. This license permits use, sharing, adaptation, distribution, and reproduction with appropriate credit, a link to the license, and an indication of any changes.",
        bold_lead="Open Access.",
        first_line_indent=False,
    )
    doc.add_heading("References", level=1)
    _add_mendeley_bibliography(doc, REFERENCES)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def build_models_zip(output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for run_dir in sorted((ROOT / "runs").glob("A09-*")):
            if not run_dir.is_dir():
                continue
            for name in (
                "config.yaml",
                "config_hash.txt",
                "reducer.joblib",
                "angle_scaler.joblib",
                "full_scaler.joblib",
                "model.joblib",
                "checkpoint_best.pt",
                "metrics.json",
                "runtime.json",
            ):
                path = run_dir / name
                if path.exists():
                    archive.write(path, path.relative_to(ROOT))
        for path in (
            ROOT / "frozen" / "protocol_v1.json",
            ROOT / "frozen" / "protocol_v1.sha256",
            ROOT / "configs" / "protocol.yaml",
            ROOT / "configs" / "selected_hparams.yaml",
            ROOT / "environment" / "pip_freeze.txt",
            ROOT / "environment" / "hardware.json",
        ):
            if path.exists():
                archive.write(path, path.relative_to(ROOT))


def _render_with_microsoft_word(docx_path: Path, output_dir: Path) -> dict:
    """Windows fallback when the bundled LibreOffice renderer is unavailable."""
    pdf_path = output_dir / f"{docx_path.stem}.pdf"
    ps_script = r"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
$document = $null
try {
    $document = $word.Documents.Open($env:ARTICLE09_DOCX, $false, $true)
    $document.ExportAsFixedFormat($env:ARTICLE09_PDF, 17)
}
finally {
    if ($null -ne $document) { $document.Close($false) }
    $word.Quit()
}
"""
    env = os.environ.copy()
    env["ARTICLE09_DOCX"] = str(docx_path.resolve())
    env["ARTICLE09_PDF"] = str(pdf_path.resolve())
    export = subprocess.run(
        ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", ps_script],
        capture_output=True,
        text=True,
        check=False,
        env=env,
    )
    if export.returncode != 0 or not pdf_path.exists():
        return {
            "returncode": export.returncode or 1,
            "stdout": export.stdout,
            "stderr": export.stderr or "Microsoft Word did not produce a PDF.",
            "rendered_pages": 0,
            "renderer": "microsoft-word",
        }
    raster_code = r"""
from pathlib import Path
import pypdfium2 as pdfium
import sys

pdf_path = Path(sys.argv[1])
out_dir = Path(sys.argv[2])
document = pdfium.PdfDocument(str(pdf_path))
for index in range(len(document)):
    page = document[index]
    image = page.render(scale=200 / 72).to_pil()
    image.save(out_dir / f"page-{index + 1}.png")
print(len(document))
"""
    raster = subprocess.run(
        [str(DOC_PYTHON), "-c", raster_code, str(pdf_path), str(output_dir)],
        capture_output=True,
        text=True,
        check=False,
    )
    pages = sorted(output_dir.glob("page-*.png"))
    return {
        "returncode": raster.returncode,
        "stdout": export.stdout + raster.stdout,
        "stderr": export.stderr + raster.stderr,
        "rendered_pages": len(pages),
        "renderer": "microsoft-word+pypdfium2",
    }


def build_manuscript_and_models(config: dict, resume: bool = True) -> None:
    notebook = ROOT / "notebooks" / "Article09_Reproducible_Research.ipynb"
    input_hash = sha256_json(
        {
            "notebook": sha256_file(notebook),
            "report": sha256_file(ROOT / "results" / "report_manifest.json"),
            "code": sha256_file(Path(__file__)),
        }
    )
    output = ROOT / "article" / "Article09_QMI_Submission.docx"
    models_zip = ROOT / "models" / "Article09_Models.zip"
    if resume and task_complete("article", input_hash) and output.exists() and models_zip.exists():
        return
    with task_context("article", input_hash):
        build_manuscript(output)
        build_models_zip(models_zip)
        qa_dir = ROOT / "article" / "qa_render"
        if not qa_dir.resolve().is_relative_to((ROOT / "article").resolve()):
            raise RuntimeError("Unsafe QA render directory")
        if qa_dir.exists():
            shutil.rmtree(qa_dir)
        qa_dir.mkdir(parents=True, exist_ok=True)
        render_command = [
            str(DOC_PYTHON),
            str(DOC_SKILL / "render_docx.py"),
            str(output),
            "--output_dir",
            str(qa_dir),
            "--emit_pdf",
        ]
        completed = subprocess.run(render_command, capture_output=True, text=True, check=False)
        pages = sorted(qa_dir.glob("page-*.png"))
        fallback = None
        if completed.returncode != 0 or not pages:
            fallback = _render_with_microsoft_word(output, qa_dir)
            pages = sorted(qa_dir.glob("page-*.png"))
        effective_returncode = fallback["returncode"] if fallback is not None else completed.returncode
        render_report = {
            "command": render_command,
            "returncode": effective_returncode,
            "stdout": completed.stdout,
            "stderr": completed.stderr,
            "rendered_pages": len(pages),
            "fallback": fallback,
            "docx_sha256": sha256_file(output),
            "models_zip_sha256": sha256_file(models_zip),
        }
        atomic_write_json(ROOT / "article" / "render_report.json", render_report)
        if effective_returncode != 0 or not pages:
            raise RuntimeError(
                "DOCX render QA failed: "
                f"returncode={effective_returncode}, rendered_pages={len(pages)}. "
                f"See {ROOT / 'article' / 'render_report.json'}."
            )

