from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUTPUT = Path(r"D:\Quantum\article10\submission_qip\QIP_Cover_Letter.docx")


def set_run(run, size=10.5, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def add_paragraph(doc, text="", *, before=0, after=7, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    if text:
        set_run(p.add_run(text))
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)

    p = add_paragraph(doc, after=2, align=WD_ALIGN_PARAGRAPH.RIGHT)
    set_run(p.add_run("16 October 2026"))

    p = add_paragraph(doc, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_run(p.add_run("Editor-in-Chief"), bold=True)
    p = add_paragraph(doc, after=10, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_run(p.add_run("Quantum Information Processing"), italic=True)

    add_paragraph(doc, "Dear Editor-in-Chief,", after=9, align=WD_ALIGN_PARAGRAPH.LEFT)

    title = "Hardware-aware simulation of quantum transfer learning across three visual domains"
    p = add_paragraph(doc, after=8)
    set_run(p.add_run("We submit the manuscript "))
    set_run(p.add_run(title), italic=True)
    set_run(p.add_run(" for consideration as a Research Article in Quantum Information Processing."))

    add_paragraph(
        doc,
        "The study evaluates a validation-selected variational quantum classifier beyond ideal statevector conditions. "
        "Using frozen DINOv2 representations and locked data partitions across CIFAR-10, PathMNIST, and EuroSAT, it "
        "quantifies finite-shot stability, controlled circuit-noise sensitivity, calibration behavior, trainability, "
        "and topology-aware transpilation overhead. The principal result is that the selected eight-qubit circuits "
        "retain predictive performance down to 128 measurement shots, while noise and restricted connectivity expose "
        "domain-specific calibration and routing costs that ideal simulation alone does not reveal.",
    )

    add_paragraph(
        doc,
        "The manuscript is appropriate for Quantum Information Processing because it connects quantum machine-learning "
        "model selection with the operational constraints that precede physical-device execution. It reports repeated "
        "seeds, validation-only selection, frozen test hashes, uncertainty intervals, finite measurement, explicit noise "
        "profiles, and coupling-map audits. The conclusions are intentionally limited to hardware-aware simulation; the "
        "manuscript does not claim quantum advantage or performance on a physical quantum processor.",
    )

    add_paragraph(
        doc,
        "This manuscript is original, has not been published, and is not under consideration by another journal. Both "
        "authors have approved the submitted version. The authors declare no competing interests and received no funding "
        "for this work. The study uses public benchmark datasets and does not involve human participants, identifiable "
        "personal data, or animal experiments. Source code, the executed notebook, configurations, checkpoints, figures, "
        "and model artifacts are available through the repository identified in the manuscript.",
    )

    add_paragraph(doc, "Thank you for considering this work.", after=13, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_paragraph(doc, "Sincerely,", after=12, align=WD_ALIGN_PARAGRAPH.LEFT)

    p = add_paragraph(doc, after=1, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_run(p.add_run("Fajri Yanuar Shiddiq Juanda"), bold=True)
    add_paragraph(doc, "Corresponding author", after=1, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_paragraph(
        doc,
        "Department of Informatics Engineering, Faculty of Computer Science",
        after=1,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_paragraph(
        doc,
        "Universitas Buana Perjuangan Karawang, Karawang, Indonesia",
        after=1,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_paragraph(doc, "Email: fajriyanuar1@gmail.com", after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_paragraph(doc, "ORCID: 0009-0007-8894-0265", after=0, align=WD_ALIGN_PARAGRAPH.LEFT)

    props = doc.core_properties
    props.title = "Cover letter for Quantum Information Processing"
    props.subject = title
    props.author = "Fajri Yanuar Shiddiq Juanda"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
